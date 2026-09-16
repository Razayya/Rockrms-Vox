#!/usr/bin/env python
"""Build the technical companion .docx that ships with a Rock Request comment.

The comment body itself is lay-readable plain text in the Outlook draft; everything an
engineer needs - ids, Guids, SQL, counts, evidence, verification steps - lives here so the
body never has to carry it. See SKILL.md in this folder.

Usage
-----
    python gen_comment_doc.py spec.json                 # writes <spec>.docx beside the json
    python gen_comment_doc.py spec.json -o out.docx

Spec (JSON)
-----------
{
  "project_id": 7607,
  "title":    "Baptism Interest split - technical notes",
  "subtitle": "Rock Request 7607",           # optional; project_id used if omitted
  "prepared": "2026-09-16",                  # optional; today if omitted
  "sections": [
    {"heading": "What this covers", "blocks": [
      {"type": "p",       "text": "..."},
      {"type": "bullets", "items": ["...", "..."]},
      {"type": "numbers", "items": ["...", "..."]},
      {"type": "kv",      "rows": [["Attribute", "4936 BaptismInterest"]]},
      {"type": "table",   "columns": ["Bucket", "People"], "rows": [["baptized", "254"]],
                          "caption": "optional"},
      {"type": "code",    "text": "SELECT ...", "caption": "optional"},
      {"type": "note",    "text": "the one thing a reader must not miss"}
    ]}
  ]
}

Every block type is optional; a section may hold any mix in any order.
"""
import argparse
import datetime as _dt
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# House palette - matches the 5720 / 6188 deliverables.
INK = RGBColor(0x36, 0x36, 0x36)
HEAD = RGBColor(0x38, 0x3B, 0x42)
MUTED = RGBColor(0x6B, 0x6B, 0x76)
GOLD_HEX = "D4AF36"
RULE_HEX = "D9D9DE"
SHADE_HEX = "F4F4F6"
NOTE_HEX = "FBF4DC"

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


# --------------------------------------------------------------------------- xml helpers
def _shade(element, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _border(element, edges, color, sz=6, val="single"):
    borders = element.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        element.append(borders)
    for edge in edges:
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)


def _spacing(par, before=0, after=6, line=None):
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    if line is not None:
        par.paragraph_format.line_spacing = line
    return par


def _run(par, text, size=10.5, bold=False, italic=False, color=INK, font=BODY_FONT):
    r = par.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    return r


# ----------------------------------------------------------------------------- doc pieces
def _title_block(doc, spec):
    par = _spacing(doc.add_paragraph(), after=2)
    _run(par, spec["title"], size=19, bold=True, color=HEAD)

    subtitle = spec.get("subtitle") or ("Rock Request " + str(spec.get("project_id", ""))).strip()
    prepared = spec.get("prepared") or _dt.date.today().isoformat()
    par = _spacing(doc.add_paragraph(), after=10)
    _run(par, subtitle + "  ·  prepared " + prepared, size=9.5, color=MUTED)
    _border(par._p.get_or_add_pPr(), ["bottom"], GOLD_HEX, sz=10)


def _heading(doc, text):
    par = _spacing(doc.add_paragraph(), before=14, after=5)
    _run(par, text, size=13, bold=True, color=HEAD)


def _paragraph(doc, text):
    par = _spacing(doc.add_paragraph(), after=7, line=1.14)
    _run(par, text)


def _list(doc, items, numbered=False):
    style = "List Number" if numbered else "List Paragraph"
    for item in items:
        par = doc.add_paragraph(style=style)
        _spacing(par, after=3, line=1.12)
        par.paragraph_format.left_indent = Inches(0.28)
        if not numbered:
            _run(par, "–  ", color=MUTED)
        _run(par, str(item))


def _grid(doc, columns, rows, widths=None):
    ncols = len(columns) if columns else (len(rows[0]) if rows else 1)
    table = doc.add_table(rows=1 if columns else 0, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    if columns:
        for cell, name in zip(table.rows[0].cells, columns):
            cell.text = ""
            par = _spacing(cell.paragraphs[0], before=3, after=3)
            _run(par, str(name), size=9.5, bold=True, color=HEAD)
            _shade(cell._tc.get_or_add_tcPr(), SHADE_HEX)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ""
            par = _spacing(cell.paragraphs[0], before=3, after=3)
            _run(par, "" if value is None else str(value), size=9.5)
    for row in table.rows:
        for cell in row.cells:
            _border(cell.paragraphs[-1]._p.get_or_add_pPr(), ["bottom"], RULE_HEX, sz=4)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    return table


def _kv(doc, rows):
    _grid(doc, [], rows, widths=[1.9, 4.6])
    _spacing(doc.add_paragraph(), after=4)


def _table(doc, block):
    _grid(doc, block.get("columns", []), block.get("rows", []))
    if block.get("caption"):
        par = _spacing(doc.add_paragraph(), before=3, after=8)
        _run(par, block["caption"], size=9, italic=True, color=MUTED)
    else:
        _spacing(doc.add_paragraph(), after=4)


def _code(doc, block):
    par = _spacing(doc.add_paragraph(), before=4, after=4, line=1.0)
    par.paragraph_format.left_indent = Inches(0.1)
    ppr = par._p.get_or_add_pPr()
    _shade(ppr, SHADE_HEX)
    _border(ppr, ["left"], GOLD_HEX, sz=12)
    lines = str(block["text"]).rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        if i:
            par.add_run().add_break()
        _run(par, line, size=8.5, font=MONO_FONT)
    if block.get("caption"):
        cap = _spacing(doc.add_paragraph(), before=2, after=8)
        _run(cap, block["caption"], size=9, italic=True, color=MUTED)


def _note(doc, block):
    par = _spacing(doc.add_paragraph(), before=6, after=8, line=1.12)
    ppr = par._p.get_or_add_pPr()
    _shade(ppr, NOTE_HEX)
    _border(ppr, ["left"], GOLD_HEX, sz=12)
    par.paragraph_format.left_indent = Inches(0.1)
    _run(par, str(block["text"]), size=10)


def _footer(doc, spec):
    par = doc.sections[0].footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    label = ("Rock Request " + str(spec.get("project_id", ""))).strip()
    _run(par, label + "  ·  " + spec["title"], size=8, color=MUTED)


RENDERERS = {
    "p": lambda doc, b: _paragraph(doc, b["text"]),
    "bullets": lambda doc, b: _list(doc, b["items"]),
    "numbers": lambda doc, b: _list(doc, b["items"], numbered=True),
    "kv": lambda doc, b: _kv(doc, b["rows"]),
    "table": _table,
    "code": _code,
    "note": _note,
}


def build(spec, out_path):
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.9)
    section.top_margin = section.bottom_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK

    _title_block(doc, spec)
    for sec in spec.get("sections", []):
        if sec.get("heading"):
            _heading(doc, sec["heading"])
        for block in sec.get("blocks", []):
            kind = block.get("type", "p")
            if kind not in RENDERERS:
                raise SystemExit("unknown block type: " + repr(kind))
            RENDERERS[kind](doc, block)
    _footer(doc, spec)

    doc.save(str(out_path))
    return out_path


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("spec", type=Path, help="JSON spec file")
    ap.add_argument("-o", "--out", type=Path, default=None)
    args = ap.parse_args(argv)

    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    for key in ("title", "sections"):
        if key not in spec:
            raise SystemExit("spec is missing required key: " + key)
    out = args.out or args.spec.with_suffix(".docx")
    build(spec, out)
    print("%s  (%s bytes, %d sections)" % (out, format(out.stat().st_size, ","), len(spec["sections"])))


if __name__ == "__main__":
    sys.exit(main())
